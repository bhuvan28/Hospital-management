from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Staff-facing headings so Excel matches the operational report language.
HEADER_LABELS = {
    "date": "Date",
    "start": "From",
    "end": "To",
    "grain": "Grain",
    "occupancy_rate": "Occupancy Rate",
    "occupied_rooms": "Occupied Rooms",
    "total_rooms": "Total Rooms",
    "average_occupancy_rate": "Average Occupancy Rate",
    "period": "Period",
    "patient_count": "Patient Count",
    "stay_count": "Stay Count",
    "open_stays": "Open Stays",
    "completed_stays": "Completed Stays",
    "new_count": "New",
    "returning_count": "Returning",
    "gender": "Gender",
    "ward": "Ward / Clinic",
    "state": "State",
    "count": "Count",
    "bucket": "Length of Stay",
    "category": "Category",
    "amount": "Amount (RM)",
    "type": "Type",
    "budgeted": "Budgeted (RM)",
    "actual": "Actual (RM)",
    "balance": "Balance (RM)",
    "vacant_rooms": "Vacant Rooms",
    "maintenance_rooms": "Maintenance Rooms",
    "todays_check_ins": "Today's Check-ins",
    "todays_check_outs": "Today's Check-outs",
    "waiting_list_open": "Waiting List",
    "total_amount": "Total Amount (RM)",
    "total_count": "Records",
    "monthly_total": "Monthly Total (RM)",
    "annual_total": "Annual Total (RM)",
    "total_budgeted": "Budgeted (RM)",
    "total_actual": "Actual (RM)",
    "total_balance": "Balance (RM)",
    "total_pax": "Pax",
    "guardian_total": "Guardians",
    "sibling_total": "Siblings",
    "family_count": "Families",
    "check_ins": "Check-ins",
    "check_outs": "Check-outs",
    "new_intakes": "New Intakes",
    "recorded_assessments": "Recorded Assessments",
    "total_assessed_fee_amount": "Assessed Amount (RM)",
    "open_entries": "Open Entries",
}


def _header_label(key: str) -> str:
    if key in HEADER_LABELS:
        return HEADER_LABELS[key]
    return key.replace("_", " ").title()


def _flatten(data: Any, prefix: str = "") -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    if isinstance(data, dict):
        for key, value in data.items():
            label = f"{prefix}{key}" if not prefix else f"{prefix}.{key}"
            if isinstance(value, (dict, list)):
                rows.extend(_flatten(value, label))
            else:
                rows.append((label, str(value)))
    elif isinstance(data, list):
        if data and isinstance(data[0], dict):
            return []
        for idx, value in enumerate(data):
            rows.extend(_flatten(value, f"{prefix}[{idx}]"))
    else:
        rows.append((prefix or "value", str(data)))
    return rows


def _table_from_report(data: dict) -> tuple[list[str], list[list[str]]]:
    for key in ("days", "trend", "groups"):
        items = data.get(key)
        if isinstance(items, list) and items and isinstance(items[0], dict):
            headers = list(items[0].keys())
            rows = [[str(row.get(h, "")) for h in headers] for row in items]
            return headers, rows
    if "by_category" in data and isinstance(data["by_category"], dict):
        first = next(iter(data["by_category"].values()), None)
        if isinstance(first, dict):
            headers = ["category", *first.keys()]
            rows = [[cat, *[str(v) for v in vals.values()]] for cat, vals in data["by_category"].items()]
            return headers, rows
        headers = ["category", "amount"]
        rows = [[k, str(v)] for k, v in data["by_category"].items()]
        return headers, rows
    if "by_type" in data:
        headers = ["type", "amount"]
        rows = [[k, str(v)] for k, v in data["by_type"].items()]
        return headers, rows
    pairs = _flatten({k: v for k, v in data.items() if not isinstance(v, (list, dict)) or k.endswith("_amount")})
    return ["field", "value"], [[_header_label(k), v] for k, v in pairs]


def to_excel_table(
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[list[Any]],
    summary: list[tuple[str, Any]] | None = None,
) -> bytes:
    """Operational workbook (room board, waitlist, audit metadata) for house staff."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    header_fill = PatternFill("solid", fgColor="0F4C5C")
    header_font = Font(color="FFFFFF", bold=True)
    ws["A1"] = "RMH HCTM UKM Operations Platform"
    ws["A1"].font = Font(bold=True, size=14, color="0F4C5C")
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = subtitle
    row_cursor = 5
    for label, value in summary or []:
        ws.cell(row=row_cursor, column=1, value=_header_label(label) if isinstance(label, str) else label)
        ws.cell(row=row_cursor, column=2, value=value)
        row_cursor += 1
    if summary:
        row_cursor += 1
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_cursor, column=col, value=_header_label(header) if header in HEADER_LABELS else header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(rows, start=row_cursor + 1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
        ws.column_dimensions[column[0].column_letter].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_excel(title: str, data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    header_fill = PatternFill("solid", fgColor="0F4C5C")
    header_font = Font(color="FFFFFF", bold=True)
    ws["A1"] = "RMH HCTM UKM Operations Platform"
    ws["A1"].font = Font(bold=True, size=14, color="0F4C5C")
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"Period: {data.get('start', '')} to {data.get('end', '')}"

    row_cursor = 5
    skip_summary = {"start", "end", "note"}
    summary_keys = [k for k, v in data.items() if k not in skip_summary and not isinstance(v, (list, dict))]
    for key in summary_keys:
        ws.cell(row=row_cursor, column=1, value=_header_label(key))
        ws.cell(row=row_cursor, column=2, value=str(data[key]))
        row_cursor += 1

    headers, rows = _table_from_report(data)
    row_cursor += 1
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_cursor, column=col, value=_header_label(header))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(rows, start=row_cursor + 1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
        ws.column_dimensions[column[0].column_letter].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(title: str, data: dict) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=title)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("RMH HCTM UKM Operations Platform", styles["Title"]),
        Paragraph(title, styles["Heading2"]),
        Paragraph(f"Period: {data.get('start', '')} to {data.get('end', '')}", styles["Normal"]),
        Spacer(1, 12),
    ]
    for key, value in data.items():
        if key in {"start", "end", "note"}:
            continue
        if not isinstance(value, (list, dict)):
            story.append(Paragraph(f"<b>{_header_label(key)}:</b> {value}", styles["Normal"]))
    story.append(Spacer(1, 12))
    headers, rows = _table_from_report(data)
    if headers:
        table_data = [[_header_label(h) for h in headers], *rows]
        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F4C5C")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.Color(0.93, 0.96, 0.96)]),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
    if data.get("note"):
        story.append(Spacer(1, 12))
        story.append(Paragraph(str(data["note"]), styles["Italic"]))
    doc.build(story)
    return buffer.getvalue()
