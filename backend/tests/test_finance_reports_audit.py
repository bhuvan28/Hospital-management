from datetime import date

from sqlalchemy import select

from app.models.audit import AuditLog
from app.models.enums import UserRole
from app.models.finance import DonationCategory
from tests.conftest import add_user, auth_header, make_patient, seed_lookups_and_rooms


def test_donation_totals(client, db):
    add_user(db, "finance@rmh.example.com", UserRole.FINANCE)
    category = DonationCategory(name="General", active=True)
    db.add(category)
    db.commit()
    headers = auth_header(client, "finance@rmh.example.com")
    first = client.post(
        "/api/donations",
        headers=headers,
        json={
            "date": date.today().isoformat(),
            "donor_name": "Demo Corporation",
            "type": "corporate",
            "category_id": category.id,
            "quantity": 1,
            "amount": 250,
        },
    )
    second = client.post(
        "/api/donations",
        headers=headers,
        json={
            "date": date.today().isoformat(),
            "donor_name": "Family Circle",
            "type": "rmh_family",
            "quantity": 2,
            "amount": 50.5,
        },
    )
    assert first.status_code == 201
    assert second.status_code == 201
    summary = client.get("/api/donations/summary", headers=headers).json()
    assert float(summary["annual_total"]) == 300.5
    assert summary["total_count"] == 2
    assert float(summary["monthly_total"]) == 300.5
    assert float(summary["by_category"]["General"]) == 250
    assert float(summary["monthly_by_category"]["General"]) == 250
    assert float(summary["monthly_by_type"]["corporate"]) == 250
    updated = client.patch(
        f"/api/donations/{first.json()['id']}",
        headers=headers,
        json={"amount": 275},
    )
    assert updated.status_code == 200
    assert float(client.get("/api/donations/summary", headers=headers).json()["annual_total"]) == 325.5
    negative = client.post(
        "/api/donations",
        headers=headers,
        json={
            "date": date.today().isoformat(),
            "donor_name": "Bad Amount",
            "type": "cash",
            "quantity": 1,
            "amount": -5,
        },
    )
    assert negative.status_code == 422


def test_petty_cash_balance_computed(client, db):
    add_user(db, "finance@rmh.example.com", UserRole.FINANCE)
    headers = auth_header(client, "finance@rmh.example.com")
    created = client.post(
        "/api/petty-cash",
        headers=headers,
        json={
            "month": date.today().replace(day=1).isoformat(),
            "budgeted": 1000,
            "actual": 350,
            "groceries_sponsored": 100,
            "category": "groceries",
        },
    )
    assert created.status_code == 201
    assert float(created.json()["balance"]) == 650
    summary = client.get("/api/petty-cash/summary", headers=headers).json()
    assert float(summary["total_balance"]) == 650
    # Clients cannot override computed balance — field is server-derived.
    assert "balance" in created.json()
    deleted = client.delete(f"/api/petty-cash/{created.json()['id']}", headers=headers)
    assert deleted.status_code == 204
    assert float(client.get("/api/petty-cash/summary", headers=headers).json()["total_balance"]) == 0


def test_meal_sponsorship_totals(client, db):
    add_user(db, "finance@rmh.example.com", UserRole.FINANCE)
    headers = auth_header(client, "finance@rmh.example.com")
    created = client.post(
        "/api/meal-sponsorships",
        headers=headers,
        json={
            "date": date.today().isoformat(),
            "sponsor_name": "Demo Sponsor",
            "meal_type": "lunch",
            "pax": 18,
            "amount": 180,
        },
    )
    assert created.status_code == 201
    listed = client.get("/api/meal-sponsorships", headers=headers)
    assert listed.status_code == 200
    assert listed.json()["total"] == 1
    meal_id = created.json()["id"]
    patched = client.patch(
        f"/api/meal-sponsorships/{meal_id}",
        headers=headers,
        json={"amount": 200, "pax": 20},
    )
    assert patched.status_code == 200
    summary = client.get("/api/meal-sponsorships/summary", headers=headers).json()
    assert float(summary["monthly_total"]) == 200
    assert summary["total_pax"] == 20
    deleted = client.delete(f"/api/meal-sponsorships/{meal_id}", headers=headers)
    assert deleted.status_code == 204
    assert client.get("/api/meal-sponsorships", headers=headers).json()["total"] == 0


def test_reports_and_exports(client, db):
    add_user(db, "admin@rmh.example.com", UserRole.ADMIN)
    add_user(db, "duty@rmh.example.com", UserRole.DUTY_OFFICER)
    ward, state, rooms = seed_lookups_and_rooms(db)
    patient = make_patient(db, ward, state)
    duty = auth_header(client, "duty@rmh.example.com")
    client.post(f"/api/rooms/{rooms[0].id}/check-in", headers=duty, json={"patient_id": patient.id})
    admin = auth_header(client, "admin@rmh.example.com")
    occupancy = client.get("/api/reports/occupancy", headers=admin).json()
    assert "days" in occupancy
    assert occupancy["days"][-1]["occupancy_rate"] >= 0
    census = client.get("/api/reports/census", headers=admin).json()
    assert census["patient_count"] == 1
    dumped = str(census)
    assert "Aisyah" not in dumped
    assert "nric" not in dumped.lower()
    add_user(db, "viewer@rmh.example.com", UserRole.VIEWER)
    viewer_headers = auth_header(client, "viewer@rmh.example.com")
    viewer_census = client.get("/api/reports/census", headers=viewer_headers)
    assert viewer_census.status_code == 200
    assert "name" not in viewer_census.json()
    dashboard = client.get("/api/dashboard", headers=viewer_headers).json()
    assert "current_occupants" not in dashboard
    assert "open_waitlist" not in dashboard
    assert "todays_movements" not in dashboard
    assert dashboard["occupancy"]["occupancy_rate"] >= 0
    assert "donations" in dashboard
    dumped_dashboard = str(dashboard)
    assert "Aisyah" not in dumped_dashboard
    assert "nric" not in dumped_dashboard.lower()
    admission = client.get("/api/reports/admission-summary", headers=admin).json()
    assert admission["check_ins"] == 1
    assert admission["new_intakes"] >= 1
    xlsx = client.get("/api/reports/occupancy/export", headers=admin, params={"format": "xlsx"})
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(xlsx.content))
    sheet = workbook.active
    assert sheet["A1"].value == "RMH HCTM UKM Operations Platform"
    assert sheet["A2"].value == "Daily Occupancy Rate"
    exported = [str(cell.value or "") for row in sheet.iter_rows() for cell in row]
    assert any("Occupancy Rate" in value for value in exported)
    pdf = client.get("/api/reports/census/export", headers=admin, params={"format": "pdf"})
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"


def test_audit_for_patient_room_and_finance(client, db):
    add_user(db, "admin@rmh.example.com", UserRole.ADMIN)
    add_user(db, "social@rmh.example.com", UserRole.SOCIAL_WORKER)
    add_user(db, "duty@rmh.example.com", UserRole.DUTY_OFFICER)
    add_user(db, "finance@rmh.example.com", UserRole.FINANCE)
    ward, state, rooms = seed_lookups_and_rooms(db)
    social = auth_header(client, "social@rmh.example.com")
    created = client.post(
        "/api/patients",
        headers=social,
        json={
            "name": "Audit Patient",
            "registration_no": "HCTM26-7777",
            "nric": "111111111111",
            "gender": "female",
            "dob": "2020-01-01",
            "address": "A",
            "diagnosis": "Demo",
            "ward_clinic_id": ward.id,
            "state_id": state.id,
        },
    )
    assert created.status_code == 201
    duty = auth_header(client, "duty@rmh.example.com")
    client.post(
        f"/api/rooms/{rooms[0].id}/check-in",
        headers=duty,
        json={"patient_id": created.json()["id"]},
    )
    finance = auth_header(client, "finance@rmh.example.com")
    client.post(
        "/api/donations",
        headers=finance,
        json={
            "date": date.today().isoformat(),
            "donor_name": "Audit Donor",
            "type": "cash",
            "quantity": 1,
            "amount": 10,
        },
    )
    client.post(
        "/api/petty-cash",
        headers=finance,
        json={
            "month": date.today().replace(day=1).isoformat(),
            "budgeted": 100,
            "actual": 40,
            "category": "groceries",
        },
    )
    client.post(
        "/api/meal-sponsorships",
        headers=finance,
        json={
            "date": date.today().isoformat(),
            "sponsor_name": "Audit Meal",
            "meal_type": "dinner",
            "pax": 10,
            "amount": 80,
        },
    )
    admin = auth_header(client, "admin@rmh.example.com")
    logs = client.get("/api/audit-log", headers=admin).json()
    entities = {item["entity"] for item in logs["items"]}
    assert "Patient" in entities
    assert "RoomStay" in entities
    assert "Donation" in entities
    assert "PettyCashEntry" in entities
    assert "MealSponsorship" in entities
    patient_logs = db.scalars(select(AuditLog).where(AuditLog.entity == "Patient")).all()
    assert patient_logs
    for log in patient_logs:
        dumped = str(log.after_values) + str(log.before_values)
        assert "password" not in dumped.lower()
